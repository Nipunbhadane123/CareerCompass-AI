import os
import re
from typing import Tuple

def extract_text_from_file(file_path: str, file_type: str) -> str:
    """Extracts raw text from PDF, DOCX, or text files safely."""
    text = ""
    file_type = file_type.lower().strip(".")
    
    if file_type == "pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            for page in doc:
                text += page.get_text("text") + "\n"
            doc.close()
        except Exception as e:
            # Fallback to plain binary read if fitz fails
            with open(file_path, "rb") as f:
                content = f.read().decode("latin-1", errors="ignore")
                text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\xff]', ' ', content)
    elif file_type in ("docx", "doc"):
        try:
            import docx
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text:
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            text = "\n".join(full_text)
        except Exception:
            with open(file_path, "rb") as f:
                text = f.read().decode("utf-8", errors="ignore")
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

    return clean_extracted_text(text)

def clean_extracted_text(text: str) -> str:
    """Sanitizes text by removing excessive whitespace and null bytes."""
    if not text:
        return ""
    text = text.replace('\x00', '')
    text = re.sub(r'\r\n', '\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()
